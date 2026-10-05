import csv
import hashlib
import json
import tempfile
import unittest
from copy import deepcopy
from importlib.resources import files
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from rowledger.cli import main
from rowledger.engine import apply_review, reconcile
from rowledger.export import export_bundle, render_html
from rowledger.rules import Columns, Rules


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.orders = self.root / "orders.csv"
        self.payments = self.root / "payments.csv"

    def write_csv(self, path, rows, headers=None):
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers or ["order_id", "amount", "currency"])
            writer.writerows(rows)

    def run_pair(self, orders, payments, rules=None):
        self.write_csv(self.orders, orders)
        self.write_csv(self.payments, payments)
        return reconcile(self.orders, self.payments, rules)

    def manual_result(self):
        return self.run_pair([["00047", "90.00", "USD"]], [["0047", "90", "USD"]])

    def review_document(self, result, **overrides):
        decision = {"order_row_id": "orders:1", "payment_row_id": "payments:1",
                    "reviewer": "Operations", "reason": "Confirmed identifier typo against source records."}
        decision.update(overrides)
        return {"schema_version": 1, "run_id": result["run_id"], "decisions": [decision]}

    def test_demo_has_independently_expected_counts_and_statuses(self):
        examples = files("rowledger").joinpath("examples")
        result = reconcile(examples.joinpath("orders.csv"), examples.joinpath("payments.csv"))
        self.assertEqual(result["summary"]["rows"], 26)
        self.assertEqual(result["summary"]["matched_pairs"], 5)
        self.assertEqual(result["summary"]["needs_attention"], 16)
        expected = {1: "matched", 4: "amount_mismatch", 5: "duplicate_key", 6: "duplicate_key",
                    7: "currency_mismatch", 8: "unmatched", 9: "invalid", 10: "invalid",
                    11: "matched", 12: "matched", 13: "duplicate_key"}
        actual = {row["record"]: row["status"] for row in result["rows"] if row["side"] == "orders"}
        for record, status in expected.items():
            self.assertEqual(actual[record], status)
        self.assertEqual(len(set(row["row_id"] for row in result["rows"])), 26)

    def test_leading_zeros_and_na_are_preserved(self):
        result = self.run_pair([["0001", "1", "USD"], ["NA", "2", "USD"]], [["1", "1", "USD"], ["NA", "2", "USD"]])
        self.assertEqual(result["rows"][0]["key"], "0001")
        self.assertEqual(result["rows"][0]["status"], "unmatched")
        self.assertEqual(result["rows"][1]["status"], "matched")

    def test_missing_identifiers_do_not_match_each_other(self):
        result = self.run_pair([["", "5", "USD"]], [["", "5", "USD"]])
        self.assertEqual([row["status"] for row in result["rows"]], ["invalid", "invalid"])

    def test_duplicate_keys_never_expand_or_guess(self):
        result = self.run_pair([["A", "1", "USD"], ["A", "2", "USD"]], [["A", "1", "USD"], ["A", "2", "USD"]])
        self.assertEqual(len(result["rows"]), 4)
        self.assertEqual({row["status"] for row in result["rows"]}, {"duplicate_key"})
        self.assertEqual(result["summary"]["matched_pairs"], 0)

    def test_invalid_duplicate_still_blocks_group(self):
        result = self.run_pair([["A", "1", "USD"], ["A", "NaN", "USD"]], [["A", "1", "USD"]])
        self.assertEqual([row["status"] for row in result["rows"]], ["duplicate_key", "invalid", "duplicate_key"])

    def test_invalid_counterpart_blocks_valid_record(self):
        result = self.run_pair([["A", "1", "USD"]], [["A", "", "USD"]])
        self.assertEqual(result["rows"][0]["status"], "blocked_invalid")

    def test_amount_and_currency_mismatches_are_distinct(self):
        result = self.run_pair([["A", "1", "USD"], ["B", "2", "USD"]], [["A", "1.01", "USD"], ["B", "2", "EUR"]])
        self.assertEqual(result["rows"][0]["status"], "amount_mismatch")
        self.assertEqual(result["rows"][1]["status"], "currency_mismatch")

    def test_decimal_equality_is_exact_and_large_amounts_are_not_floats(self):
        result = self.run_pair([["A", "9007199254740993.10", "USD"], ["B", "0.3000", "USD"]], [["A", "9007199254740993.1", "USD"], ["B", "0.30", "USD"]])
        self.assertEqual(result["summary"]["matched_pairs"], 2)
        self.assertEqual(result["rows"][0]["amount_minor"], "900719925474099310")

    def test_invalid_money_is_quarantined_without_rounding(self):
        values = ["NaN", "Infinity", "1e2", "-1", "1,000", "1.001", "１２", "", ".5"]
        result = self.run_pair([[str(index), value, "USD"] for index, value in enumerate(values)], [])
        self.assertTrue(all(row["status"] == "invalid" for row in result["rows"]))

    def test_currency_scale_is_explicit(self):
        result = self.run_pair([["A", "1.00", "TWD"], ["B", "1.5", "TWD"], ["C", "1", "usd"]], [["A", "1", "TWD"]])
        self.assertEqual([row["status"] for row in result["rows"][:3]], ["matched", "invalid", "invalid"])

    def test_trimming_is_configurable_and_raw_value_stays_intact(self):
        result = self.run_pair([[" A ", "1", "USD"]], [["A", "1", "USD"]])
        self.assertEqual(result["rows"][0]["status"], "matched")
        self.assertEqual(result["rows"][0]["raw"]["order_id"], " A ")
        strict = reconcile(self.orders, self.payments, Rules(trim_keys=False))
        self.assertEqual(strict["summary"]["matched_pairs"], 0)

    def test_inputs_remain_byte_identical_and_reruns_are_deterministic(self):
        result = self.manual_result()
        before = self.orders.read_bytes()
        second = reconcile(self.orders, self.payments)
        self.assertEqual(result, second)
        self.assertEqual(self.orders.read_bytes(), before)
        self.assertEqual(result["sources"]["orders"]["sha256"], hashlib.sha256(before).hexdigest())

    def test_manual_pair_keeps_original_keys_and_base_result(self):
        result = self.manual_result()
        reviewed = apply_review(result, self.review_document(result))
        self.assertEqual(reviewed["summary"]["matched_pairs"], 1)
        self.assertEqual([row["key"] for row in reviewed["rows"]], ["00047", "0047"])
        self.assertEqual(result["summary"]["matched_pairs"], 0)

    def test_stale_review_rejects_changed_input(self):
        result = self.manual_result()
        document = self.review_document(result)
        self.write_csv(self.orders, [["00047", "91", "USD"]])
        with self.assertRaisesRegex(ValueError, "Stale"):
            apply_review(reconcile(self.orders, self.payments), document)

    def test_stale_review_rejects_changed_rules_or_filename(self):
        result = self.manual_result()
        document = self.review_document(result)
        with self.assertRaisesRegex(ValueError, "Stale"):
            apply_review(reconcile(self.orders, self.payments, Rules(trim_keys=False)), document)
        renamed = self.root / "new-orders.csv"
        renamed.write_bytes(self.orders.read_bytes())
        with self.assertRaisesRegex(ValueError, "Stale"):
            apply_review(reconcile(renamed, self.payments), document)

    def test_review_cannot_reuse_a_row(self):
        result = self.manual_result()
        document = self.review_document(result)
        document["decisions"] *= 2
        with self.assertRaisesRegex(ValueError, "two review decisions"):
            apply_review(result, document)

    def test_review_cannot_override_amount_currency_or_invalid_data(self):
        for payment in (["B", "91", "USD"], ["B", "90", "EUR"], ["B", "bad", "USD"]):
            with self.subTest(payment=payment):
                result = self.run_pair([["A", "90", "USD"]], [payment])
                with self.assertRaises(ValueError):
                    apply_review(result, self.review_document(result))

    def test_review_cannot_replace_automatic_match(self):
        result = self.run_pair([["A", "1", "USD"]], [["A", "1", "USD"]])
        with self.assertRaisesRegex(ValueError, "already matched"):
            apply_review(result, self.review_document(result))

    def test_review_requires_reason_known_ids_and_correct_sides(self):
        result = self.manual_result()
        for override in ({"reason": " "}, {"reviewer": ""}, {"order_row_id": "no-row"},
                         {"order_row_id": "payments:1"}, {"reason": "x" * 1001}, {"reason": "A\x00B"}):
            with self.subTest(override=override), self.assertRaises(ValueError):
                apply_review(result, self.review_document(result, **override))

    def test_review_schema_is_strict(self):
        result = self.manual_result()
        document = self.review_document(result)
        for change in ({"schema_version": True}, {"decisions": {}}, {"extra": 1}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                apply_review(result, {**document, **change})

    def test_failed_review_is_atomic(self):
        result = self.run_pair([["A", "1", "USD"], ["B", "2", "USD"]], [["X", "1", "USD"], ["Y", "3", "USD"]])
        before = deepcopy(result)
        document = self.review_document(result)
        document["decisions"].append({**document["decisions"][0], "order_row_id": "orders:2", "payment_row_id": "payments:2"})
        with self.assertRaises(ValueError):
            apply_review(result, document)
        self.assertEqual(result, before)

    def test_duplicate_and_malformed_headers_are_rejected(self):
        self.write_csv(self.payments, [])
        for headers in (["order_id", "amount", "amount"], ["order_id", "", "currency"], ["wrong", "amount", "currency"]):
            self.write_csv(self.orders, [], headers)
            with self.subTest(headers=headers), self.assertRaises(ValueError):
                reconcile(self.orders, self.payments)

    def test_bad_csv_rows_and_blank_lines_are_not_silently_dropped(self):
        self.write_csv(self.payments, [])
        for content in ("order_id,amount,currency\nA,1\n", "order_id,amount,currency\n\n", 'order_id,amount,currency\n"unfinished'):
            self.orders.write_text(content)
            with self.subTest(content=content), self.assertRaises((ValueError, csv.Error)):
                reconcile(self.orders, self.payments)

    def test_csv_record_reference_handles_multiline_fields(self):
        result = self.run_pair([["A\nB", "1", "USD"], ["C", "2", "USD"]], [])
        self.assertEqual(result["rows"][1]["location"], "record 2")
        self.assertEqual(result["rows"][0]["raw"]["order_id"], "A\nB")

    def test_header_only_inputs_are_supported(self):
        result = self.run_pair([], [])
        self.assertEqual(result["summary"]["rows"], 0)
        export_bundle(result, self.root / "empty")
        self.assertIsNone(result["profiles"]["orders"]["empty_ratio"]["order_id"])

    def test_control_characters_are_rejected_instead_of_stripped(self):
        for value in ("A\x00B", "A\uffffB"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "control characters"):
                self.run_pair([[value, "1", "USD"]], [])

    def test_file_row_and_cell_limits_fail_closed(self):
        self.run_pair([["A", "1", "USD"], ["B", "2", "USD"]], [])
        with patch("rowledger.readers.MAX_ROWS", 1), self.assertRaises(ValueError):
            reconcile(self.orders, self.payments)
        with patch("rowledger.readers.MAX_BYTES", 5), self.assertRaises(ValueError):
            reconcile(self.orders, self.payments)
        self.write_csv(self.orders, [["A" * 32768, "1", "USD"]])
        with self.assertRaises(ValueError):
            reconcile(self.orders, self.payments)

    def write_xlsx(self, rows, second_sheet=False):
        path = self.root / "orders.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Orders"
        worksheet.append(["order_id", "amount", "currency"])
        for row in rows:
            worksheet.append(row)
        if second_sheet:
            workbook.create_sheet("Notes")
        workbook.save(path)
        return path

    def test_xlsx_numeric_identifier_and_formula_are_flagged(self):
        self.write_csv(self.payments, [["0001", "1", "USD"]])
        path = self.write_xlsx([[1, 1, "USD"], ["0002", "=1+1", "USD"]])
        result = reconcile(path, self.payments)
        self.assertEqual([row["status"] for row in result["rows"][:2]], ["invalid", "invalid"])

    def test_xlsx_text_identifier_numeric_amount_and_row_provenance(self):
        self.write_csv(self.payments, [["0001", "1.25", "USD"]])
        path = self.write_xlsx([["0001", 1.25, "USD"]])
        result = reconcile(path, self.payments)
        self.assertEqual(result["summary"]["matched_pairs"], 1)
        self.assertEqual(result["rows"][0]["location"], "row 2")
        self.assertEqual(result["rows"][0]["sheet"], "Orders")

    def test_xlsx_multisheet_requires_explicit_mapping(self):
        self.write_csv(self.payments, [])
        path = self.write_xlsx([["A", 1, "USD"]], second_sheet=True)
        with self.assertRaisesRegex(ValueError, "explicitly"):
            reconcile(path, self.payments)
        result = reconcile(path, self.payments, Rules(orders=Columns(sheet="Orders")))
        self.assertEqual(len(result["rows"]), 1)

    def test_xlsx_expansion_limit_is_enforced(self):
        self.write_csv(self.payments, [])
        path = self.write_xlsx([["A", 1, "USD"]])
        with patch("rowledger.readers.MAX_EXPANDED_BYTES", 1), self.assertRaises(ValueError):
            reconcile(path, self.payments)

    def test_custom_mapping_and_unknown_rule_fields(self):
        self.write_csv(self.orders, [["A", "1", "USD"]], ["ID", "Total", "Unit"])
        self.write_csv(self.payments, [["A", "1", "USD"]])
        result = reconcile(self.orders, self.payments, Rules.from_dict({"orders": {"key": "ID", "amount": "Total", "currency": "Unit"}}))
        self.assertEqual(result["summary"]["matched_pairs"], 1)
        for data in ({"typo": 1}, {"minor_units": {"USD": True}}, {"trim_keys": "true"},
                     {"orders": {"key": "amount"}}, {"payments": {"missing": 1}}):
            with self.subTest(data=data), self.assertRaises(ValueError):
                Rules.from_dict(data)

    def test_export_is_literal_and_html_payload_is_escaped(self):
        payload = '</script><script>alert("injected")</script>'
        result = self.run_pair([["=1+1", "1", "USD"], [payload, "2", "USD"]], [])
        output = export_bundle(result, self.root / "run")
        html = (output / "report.html").read_text()
        self.assertNotIn(payload, html)
        self.assertIn("\\u003c/script\\u003e", html)
        self.assertIn("connect-src 'none'", html)
        workbook = load_workbook(output / "results.xlsx", data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["orders_input"]["B2"].value, "=1+1")
        self.assertEqual(workbook["orders_input"]["B2"].data_type, "s")
        manifest = json.loads((output / "manifest.json").read_text())
        for name, digest in manifest["sha256"].items():
            self.assertEqual(hashlib.sha256((output / name).read_bytes()).hexdigest(), digest)

    def test_existing_output_directory_is_preserved(self):
        result = self.manual_result()
        output = self.root / "existing"
        output.mkdir()
        marker = output / "keep.txt"
        marker.write_text("preserve")
        with self.assertRaisesRegex(ValueError, "already exists"):
            export_bundle(result, output)
        self.assertEqual(marker.read_text(), "preserve")

    def test_cli_demo_and_review_round_trip(self):
        output = self.root / "demo"
        self.assertEqual(main(["demo", "--out", str(output)]), 0)
        result = json.loads((output / "results.json").read_text())
        document = {"schema_version": 1, "run_id": result["run_id"], "decisions": [
            {"order_row_id": "orders:8", "payment_row_id": "payments:7", "reviewer": "Operations",
             "reason": "Synthetic example: leading-zero typo confirmed."}]}
        review = self.root / "review.json"
        review.write_text(json.dumps(document))
        final = self.root / "reviewed"
        self.assertEqual(main(["run", str(output / "inputs/orders.csv"), str(output / "inputs/payments.csv"),
                               "--review", str(review), "--out", str(final)]), 0)
        reviewed = json.loads((final / "results.json").read_text())
        self.assertEqual(reviewed["summary"]["matched_pairs"], 6)
        self.assertEqual(reviewed["summary"]["needs_attention"], 14)


if __name__ == "__main__":
    unittest.main()
