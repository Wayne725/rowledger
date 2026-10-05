import hashlib
import json
import tempfile
from importlib.resources import files
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from .engine import SCHEMA_VERSION


def _literal_row(worksheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(worksheet, value=str(value) if value is not None else "")
        cell.data_type = "s"
        cells.append(cell)
    worksheet.append(cells)


def write_xlsx(result, path):
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("results")
    worksheet.freeze_panes = "A2"
    columns = ["row_id", "side", "source", "sheet", "location", "key", "currency", "amount",
               "amount_minor", "status", "partner_id", "explanation", "reviewer", "reason"]
    _literal_row(worksheet, columns)
    for row in result["rows"]:
        review = row["review"] or {}
        _literal_row(worksheet, [review.get(column, row.get(column)) for column in columns])
    for side in ("orders", "payments"):
        worksheet = workbook.create_sheet(f"{side}_input")
        worksheet.freeze_panes = "A2"
        headers = result["sources"][side]["columns"]
        _literal_row(worksheet, ["source_record", *headers])
        for row in result["rows"]:
            if row["side"] == side:
                _literal_row(worksheet, [row["record"], *[row["raw"][column] for column in headers]])
    workbook.save(path)


def render_html(result):
    assets = files("rowledger").joinpath("assets")
    payload = json.dumps(result, ensure_ascii=False, allow_nan=False).replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e")
    return (assets.joinpath("report.html").read_text()
            .replace("__STYLE__", assets.joinpath("report.css").read_text())
            .replace("__SCRIPT__", assets.joinpath("report.js").read_text())
            .replace("__RESULT__", payload))


def export_bundle(result, output_dir):
    output_dir = Path(output_dir)
    if output_dir.exists():
        raise ValueError("Output directory already exists. Choose a new directory to preserve earlier runs.")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".rowledger-", dir=output_dir.parent) as staging:
        stage = Path(staging)
        (stage / "results.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        (stage / "report.html").write_text(render_html(result), encoding="utf-8")
        template = {"schema_version": SCHEMA_VERSION, "run_id": result["run_id"], "decisions": result["decisions"]}
        (stage / "review-template.json").write_text(json.dumps(template, ensure_ascii=False, indent=2), encoding="utf-8")
        write_xlsx(result, stage / "results.xlsx")
        checksums = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(stage.iterdir())}
        (stage / "manifest.json").write_text(json.dumps({"run_id": result["run_id"], "sha256": checksums}, indent=2), encoding="utf-8")
        stage.rename(output_dir)
    return output_dir
