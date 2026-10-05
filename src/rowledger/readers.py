import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .rules import Columns

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 80 * 1024 * 1024
MAX_ROWS = 25_000
MAX_COLUMNS = 100


@dataclass
class Source:
    name: str
    sha256: str
    sheet: str | None
    frame: pd.DataFrame
    cell_errors: dict[int, list[str]]

    def metadata(self):
        return {"name": self.name, "sha256": self.sha256, "sheet": self.sheet,
                "records": len(self.frame), "columns": self.frame.columns.tolist()}

    def profile(self):
        self.frame.head()
        info = io.StringIO()
        self.frame.info(buf=info)
        description = self.frame.describe(include="all").fillna("") if len(self.frame) else pd.DataFrame()
        empty_ratio = {column: float(value) if pd.notna(value) else None
                       for column, value in self.frame.eq("").mean().items()}
        return {"info": info.getvalue(), "empty_ratio": empty_ratio,
                "unique_values": {column: int(self.frame[column].nunique()) for column in self.frame},
                "described_columns": description.columns.tolist()}


def _headers(values):
    if not values or len(values) > MAX_COLUMNS:
        raise ValueError(f"Expected 1–{MAX_COLUMNS} columns.")
    if any(not isinstance(value, str) or not value.strip() or len(value) > 256 for value in values):
        raise ValueError("Headers must be nonempty text.")
    if len(set(values)) != len(values):
        raise ValueError("Duplicate column headers are not allowed.")
    return values


def read_source(path: Path | str, columns: Columns) -> Source:
    path = Path(path)
    with path.open("rb") as handle:
        content = handle.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValueError("Input exceeds the 10 MiB limit.")
    errors = {}
    sheet = None
    if path.suffix.lower() == ".csv":
        if columns.sheet is not None:
            raise ValueError("A CSV input cannot have a sheet mapping.")
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        headers = _headers(next(reader, []))
        records = []
        for record_no, values in enumerate(reader, 1):
            if record_no > MAX_ROWS:
                raise ValueError(f"Input exceeds {MAX_ROWS} data records.")
            if len(values) != len(headers):
                raise ValueError(f"CSV record {record_no} has {len(values)} fields; expected {len(headers)}.")
            if any(len(value) > 32_767 for value in values):
                raise ValueError("A field exceeds the XLSX text-cell limit of 32,767 characters.")
            records.append(values)
    elif path.suffix.lower() == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 2_000 or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
                raise ValueError("XLSX archive exceeds the expansion limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if columns.sheet is None and len(workbook.sheetnames) != 1:
                raise ValueError("Choose a sheet explicitly for a workbook with multiple sheets.")
            sheet = columns.sheet or workbook.sheetnames[0]
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Sheet {sheet!r} does not exist.")
            worksheet = workbook[sheet]
            worksheet.reset_dimensions()
            iterator = worksheet.iter_rows()
            header_cells = next(iterator, ())
            headers = _headers([cell.value for cell in header_cells])
            records = []
            for record_no, cells in enumerate(iterator, 1):
                if record_no > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise ValueError("XLSX exceeds the row or column limit.")
                if len(cells) > len(headers):
                    raise ValueError("XLSX data extends past the header columns.")
                values = ["" if cell.value is None else str(cell.value) for cell in cells]
                values += [""] * (len(headers) - len(values))
                row_errors = []
                for index, cell in enumerate(cells):
                    if headers[index] in (columns.key, columns.amount, columns.currency):
                        if cell.data_type in ("f", "e"):
                            row_errors.append(f"{headers[index]} contains a formula or Excel error.")
                        elif headers[index] == columns.key and cell.value is not None and not isinstance(cell.value, str):
                            row_errors.append("Identifier is not stored as text; lost leading zeros cannot be inferred.")
                if row_errors:
                    errors[record_no] = row_errors
                records.append(values)
        finally:
            workbook.close()
    else:
        raise ValueError("Supported inputs are UTF-8 CSV and XLSX.")
    missing = set((columns.key, columns.amount, columns.currency)) - set(headers)
    if missing:
        raise ValueError(f"Missing columns: {', '.join(sorted(missing))}.")
    if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", value)
           for values in [headers, *records] for value in values):
        raise ValueError("Input contains control characters that cannot be preserved in XLSX.")
    frame = pd.DataFrame(records, columns=headers, dtype=str)
    return Source(path.name, hashlib.sha256(content).hexdigest(), sheet, frame, errors)
